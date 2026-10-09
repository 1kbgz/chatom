import io
from zipfile import ZIP_DEFLATED, ZipFile

import pytest


@pytest.mark.parametrize("filename", ["report.txt", "report.csv", "report.xlsx", "report.pdf", "report.docx"])
def test_create_and_read_document(filename):
    from chatom.agent.files import create_document, read_document

    data, media_type = create_document(filename, text="Quarterly revenue", columns=["Team", "Revenue"], rows=[["Alpha", 42]])
    result = read_document(data, filename, media_type)
    assert "Quarterly revenue" in str(result) or "Alpha" in str(result)
    if filename.endswith((".csv", ".xlsx", ".docx")):
        assert "Alpha" in str(result) and "42" in str(result)
    if filename.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data))
        assert workbook.active.cell(2, 2).value == 42


def test_document_reader_rejects_unsupported_format():
    from chatom.agent.files import read_document

    with pytest.raises(ValueError, match="Unsupported"):
        read_document(b"binary", "file.exe", "application/octet-stream")


def test_document_reader_rejects_expanding_archive():
    from chatom.agent.files import read_document

    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * 1_000_001)
    with pytest.raises(ValueError, match="expanded"):
        read_document(output.getvalue(), "report.xlsx", max_bytes=1_000_000)


def test_spreadsheet_text_is_not_executed_as_formula():
    from openpyxl import load_workbook

    from chatom.agent.files import create_document

    data, _ = create_document("report.xlsx", columns=["Value"], rows=[["=1+1"]])
    workbook = load_workbook(io.BytesIO(data))
    assert workbook.active.cell(2, 1).data_type == "s"


def test_document_reader_reports_text_truncation():
    from chatom.agent.files import read_document

    result = read_document(b"abcdef", "report.txt", max_chars=3)
    assert result["text"] == "abc" and result["truncated"] is True


@pytest.mark.asyncio
async def test_document_tool_creates_and_uploads_workbook():
    from chatom.agent.toolset import BackendToolset
    from chatom.tests.test_agent import _MockBackend

    backend = _MockBackend()
    result = await BackendToolset(backend).call(
        "create_file", {"channel": {"id": "C1"}, "filename": "report.xlsx", "columns": ["Value"], "rows": [[42]]}
    )
    assert result["ok"] is True and result["message_id"]
    from openpyxl import load_workbook

    assert load_workbook(io.BytesIO(backend.uploaded[0]["data"])).active.cell(2, 1).value == 42


@pytest.mark.asyncio
async def test_image_tool_generates_and_uploads_without_model_base64():
    from unittest.mock import AsyncMock

    from chatom.agent.toolset import BackendToolset
    from chatom.tests.test_agent import _MockBackend

    backend = _MockBackend()
    generate = AsyncMock(return_value=b"\x89PNG\r\n\x1a\nimage")
    result = await BackendToolset(backend, image_generator=generate).call("generate_image", {"channel": {"id": "C1"}, "prompt": "A blue square"})
    assert result["ok"] and "data_base64" not in result
    generate.assert_awaited_once_with("A blue square")
    assert backend.uploaded[0]["filename"].endswith(".png")


def test_read_only_toolset_does_not_offer_file_writers():
    from chatom.agent.toolset import BackendToolset
    from chatom.tests.test_agent import _MockBackend

    toolset = BackendToolset(_MockBackend(), read_only=True, image_generator=lambda prompt: None)
    assert {"create_file", "generate_image"}.isdisjoint(toolset.tool_definitions())
