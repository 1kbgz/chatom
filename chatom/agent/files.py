from __future__ import annotations

import csv
import io
import mimetypes
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape
from zipfile import ZipFile

MAX_BYTES = 5_000_000
MAX_CHARS = 50_000
MAX_CELLS = 10_000
MAX_PAGES = 100


def _format(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in {".txt", ".csv", ".xlsx", ".pdf", ".docx"}:
        raise ValueError("Unsupported document format. Use TXT, CSV, XLSX, PDF, or DOCX.")
    return extension


def _check_table(columns: list[str], rows: list[list[Any]]) -> None:
    if len(columns) + sum(len(row) for row in rows) > MAX_CELLS:
        raise ValueError("Document exceeds the cell limit.")
    if columns and any(len(row) != len(columns) for row in rows):
        raise ValueError("Every row must match the number of columns.")


def create_document(
    filename: str,
    *,
    text: str = "",
    columns: list[str] | None = None,
    rows: list[list[Any]] | None = None,
    sheet_name: str = "Data",
) -> tuple[bytes, str]:
    extension = _format(filename)
    columns, rows = columns or [], rows or []
    _check_table(columns, rows)
    if len(text) + sum(len(str(value)) for row in [columns, *rows] for value in row) > MAX_CHARS:
        raise ValueError("Document exceeds the text limit.")
    output = io.BytesIO()
    if extension == ".txt":
        data = text.encode("utf-8")
    elif extension == ".csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        for row in ([columns] if columns else []) + rows:
            writer.writerow(["'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value for value in row])
        data = stream.getvalue().encode("utf-8")
    elif extension == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        for row in ([columns] if columns else []) + rows:
            worksheet.append(row)
            for cell in worksheet[worksheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
    elif extension == ".docx":
        from docx import Document

        document = Document()
        document.add_paragraph(text)
        table_rows = ([columns] if columns else []) + rows
        if table_rows:
            table = document.add_table(rows=0, cols=max(len(row) for row in table_rows))
            for row in table_rows:
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = str(value)
        document.save(output)
        data = output.getvalue()
    else:
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Table

        styles = getSampleStyleSheet()
        items = [Paragraph(escape(text).replace("\n", "<br/>"), styles["BodyText"])]
        table_rows = ([columns] if columns else []) + rows
        if table_rows:
            items.append(Table([[Paragraph(escape(str(value)), styles["BodyText"]) for value in row] for row in table_rows]))
        SimpleDocTemplate(output).build(items)
        data = output.getvalue()
    if len(data) > MAX_BYTES:
        raise ValueError("Document exceeds the byte limit.")
    office_types = {
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    return data, office_types.get(extension) or mimetypes.guess_type(filename)[0] or "application/octet-stream"


def read_document(
    data: bytes,
    filename: str,
    content_type: str = "",
    *,
    max_bytes: int = MAX_BYTES,
    max_chars: int = MAX_CHARS,
) -> dict[str, Any]:
    extension = _format(filename)
    if len(data) > max_bytes or max_bytes < 1 or max_chars < 1:
        raise ValueError("Document exceeds the byte limit or has invalid limits.")
    if extension in {".xlsx", ".docx"}:
        with ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 2_000 or sum(item.file_size for item in archive.infolist()) > max_bytes:
                raise ValueError("Document expanded archive exceeds the byte limit.")
    result: dict[str, Any] = {"filename": filename, "content_type": content_type or mimetypes.guess_type(filename)[0], "truncated": False}
    text = ""
    if extension == ".txt":
        text = data.decode("utf-8-sig")
    elif extension == ".csv":
        table = []
        cells = 0
        for row in csv.reader(io.StringIO(data.decode("utf-8-sig"))):
            cells += len(row)
            if cells > MAX_CELLS or sum(len(str(value)) for value in row) + len(text) > max_chars:
                result["truncated"] = True
                break
            table.append(row)
            text += "\t".join(row) + "\n"
        result["rows"] = table
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheets: dict[str, list[list[Any]]] = {}
        cells = 0
        try:
            for sheet in workbook:
                values = []
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    row_text = "\t".join(str(value) if value is not None else "" for value in row) + "\n"
                    if cells > MAX_CELLS or len(text) + len(row_text) > max_chars:
                        result["truncated"] = True
                        break
                    values.append([value.isoformat() if hasattr(value, "isoformat") else value for value in row])
                    text += row_text
                sheets[sheet.title] = values
                if result["truncated"]:
                    break
        finally:
            workbook.close()
        result["sheets"] = sheets
    elif extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + "\t".join(cell.text for cell in row.cells)
                if len(text) > max_chars:
                    break
            if len(text) > max_chars:
                break
    else:
        from pypdf import PdfReader

        document = PdfReader(io.BytesIO(data))
        if document.is_encrypted:
            raise ValueError("Encrypted PDF documents are unsupported.")
        result["pages"] = len(document.pages)
        image_pages = []
        for index, page in enumerate(document.pages[:MAX_PAGES]):
            page_text = page.extract_text() or ""
            if not page_text.strip():
                image_pages.append(index + 1)
            text += page_text + "\n"
            if len(text) > max_chars:
                break
        result["needs_ocr_pages"] = image_pages
        result["truncated"] = len(document.pages) > MAX_PAGES
    result["text"] = text[:max_chars]
    result["truncated"] = result["truncated"] or len(text) > max_chars
    return result
